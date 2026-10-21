"""Turn uploaded documents into Markdown text a model can read.

Optional module: ``pip install "cairnkit[documents]"`` adds pypdf (PDF) and openpyxl
(Excel). Word (.docx) and plain text need only the standard library.

    from cairnkit.documents import DocumentError, extract_text

    try:
        doc = extract_text(data, "price-list.xlsx")
    except DocumentError as exc:
        show(exc.code)  # stable code for the host to translate
    else:
        use(doc.text)

The text is untrusted data. Hosts must keep it out of system prompts and tell the model
that instructions inside a document are not instructions to it.

Backends are replaceable: ``register_extractor("pdf", my_pdf_reader)`` swaps in, for
example, MarkItDown or an OCR service for one format without changing callers.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime, time
from xml.etree import ElementTree

__all__ = [
    "SUPPORTED_EXTENSIONS",
    "Document",
    "DocumentError",
    "Limits",
    "Extracted",
    "detect_kind",
    "extract_text",
    "register_extractor",
]

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".xlsx", ".txt", ".md", ".csv", ".tsv")
_TEXT_EXTENSIONS = (".txt", ".md", ".markdown", ".csv", ".tsv", ".text")


class DocumentError(Exception):
    """A document that cannot be turned into text. ``code`` is stable; ``str()`` is English.

    Codes: ``too_large``, ``empty``, ``unsupported``, ``old_format`` (.doc/.xls/.ppt),
    ``image`` (needs a vision model), ``encrypted``, ``no_text`` (scanned PDF),
    ``unreadable`` (damaged or not what it claims), ``needs_extra`` (missing optional
    dependency).
    """

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class Limits:
    max_bytes: int = 20 * 1024 * 1024  # the upload itself
    max_unzipped_bytes: int = 200 * 1024 * 1024  # .docx/.xlsx are zip files: no zip bombs
    max_zip_members: int = 10_000
    max_pages: int = 300  # PDF pages read
    max_sheet_rows: int = 5_000  # rows per Excel sheet
    max_chars: int = 200_000  # the Markdown handed back


@dataclass(frozen=True)
class Extracted:
    """What an extractor returns; ``extract_text`` adds the common clean-up and limits."""

    text: str
    parts: int | None = None  # PDF pages or Excel sheets read
    warnings: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class Document:
    text: str  # Markdown
    kind: str  # "pdf" | "docx" | "xlsx" | "text"
    parts: int | None  # PDF pages / Excel sheets, when meaningful
    truncated: bool
    warnings: tuple[str, ...]

    @property
    def chars(self) -> int:
        return len(self.text)


Extractor = Callable[[bytes, Limits], Extracted]
_EXTRACTORS: dict[str, Extractor] = {}


def register_extractor(kind: str, extractor: Extractor) -> Extractor | None:
    """Use ``extractor`` for ``kind`` ("pdf", "docx", "xlsx", "text"); returns the old one."""
    previous = _EXTRACTORS.get(kind)
    _EXTRACTORS[kind] = extractor
    return previous


# ── detection ────────────────────────────────────────────────────────────────

_IMAGE_MAGIC = (b"\x89PNG", b"\xff\xd8\xff", b"GIF8", b"BM")


def detect_kind(data: bytes, filename: str = "") -> str:
    """The document kind, judged by content first and the file name second."""
    name = filename.lower()
    head = data[:2048]
    if b"%PDF-" in head[:1024]:
        return "pdf"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        raise DocumentError(
            "old_format",
            "This is an old Office file (.doc/.xls/.ppt). Save it as .docx or .xlsx and retry.",
        )
    if head.startswith(_IMAGE_MAGIC) or (head[:4] == b"RIFF" and head[8:12] == b"WEBP"):
        raise DocumentError("image", "This is a picture, not a document.")
    if head.startswith(b"PK\x03\x04"):
        try:
            names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
        except zipfile.BadZipFile as exc:
            raise DocumentError("unreadable", "The file is damaged.") from exc
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        raise DocumentError("unsupported", "Only PDF, Word (.docx), Excel (.xlsx) and text.")
    if name.endswith(_TEXT_EXTENSIONS) or _looks_like_text(head):
        return "text"
    raise DocumentError("unsupported", "Only PDF, Word (.docx), Excel (.xlsx) and text.")


def _looks_like_text(head: bytes) -> bool:
    if b"\x00" in head:
        return False
    try:
        head.decode("utf-8")
    except UnicodeDecodeError as exc:
        return exc.start >= len(head) - 3  # a character cut at the end of the sample
    return True


# ── entry point ──────────────────────────────────────────────────────────────


def extract_text(data: bytes, filename: str = "", *, limits: Limits | None = None) -> Document:
    """Markdown text of a PDF, Word, Excel or text file. Raises :class:`DocumentError`."""
    limits = limits or Limits()
    if len(data) > limits.max_bytes:
        raise DocumentError("too_large", f"Files can be at most {limits.max_bytes // 2**20} MB.")
    if not data:
        raise DocumentError("empty", "The file is empty.")
    kind = detect_kind(data, filename)
    if kind in ("docx", "xlsx"):
        _check_zip(data, limits)
    try:
        result = _EXTRACTORS[kind](data, limits)
    except DocumentError:
        raise
    except Exception as exc:  # parsers raise many kinds of errors on damaged files
        raise DocumentError("unreadable", "The file could not be read.") from exc
    text = _tidy(result.text)
    if not text:
        raise DocumentError("no_text", "No text was found in this file.")
    warnings = list(result.warnings)
    truncated = len(text) > limits.max_chars
    if truncated:
        cut = text.rfind("\n", 0, limits.max_chars)
        text = text[: cut if cut > limits.max_chars // 2 else limits.max_chars].rstrip()
        warnings.append(f"Only the first {limits.max_chars:,} characters were kept.")
    return Document(text, kind, result.parts, truncated, tuple(warnings))


def _check_zip(data: bytes, limits: Limits) -> None:
    infos = zipfile.ZipFile(io.BytesIO(data)).infolist()
    if len(infos) > limits.max_zip_members or (
        sum(i.file_size for i in infos) > limits.max_unzipped_bytes
    ):
        raise DocumentError("too_large", "The file expands to too much data to read safely.")


_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_BLANK_RUNS = re.compile(r"\n{3,}")


def _tidy(text: str) -> str:
    text = _CONTROL.sub("", text.replace("\r\n", "\n").replace("\r", "\n"))
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return _BLANK_RUNS.sub("\n\n", text).strip()


# ── plain text ───────────────────────────────────────────────────────────────


def _text(data: bytes, limits: Limits) -> Extracted:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return Extracted(data.decode("utf-16", errors="replace"))
    try:
        return Extracted(data.decode("utf-8-sig"))
    except UnicodeDecodeError:
        # Windows-1252 decodes almost anything; say so, as some characters may be off.
        return Extracted(
            data.decode("cp1252", errors="replace"), warnings=["The text was not UTF-8."]
        )


# ── PDF (pypdf) ──────────────────────────────────────────────────────────────


def _pdf(data: bytes, limits: Limits) -> Extracted:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentError(
            "needs_extra", 'Reading PDF needs: pip install "cairnkit[documents]"'
        ) from exc
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            opened = reader.decrypt("")  # many PDFs only restrict printing or copying
        except Exception as exc:
            raise DocumentError("encrypted", "This PDF is password protected.") from exc
        if not opened:
            raise DocumentError("encrypted", "This PDF is password protected.")
    total = len(reader.pages)
    pages, blank = [], 0
    for number, page in enumerate(reader.pages[: limits.max_pages], start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            blank += 1
            continue
        pages.append(f"[Page {number}]\n{text}")
    warnings = []
    read = min(total, limits.max_pages)
    if not pages:
        raise DocumentError(
            "no_text", "This PDF has no selectable text; it is probably a scan or photos."
        )
    if blank:
        warnings.append(f"{blank} of {read} pages had no text (scanned pages are skipped).")
    if total > limits.max_pages:
        warnings.append(f"Only the first {limits.max_pages} of {total} pages were read.")
    return Extracted("\n\n".join(pages), read, warnings)


# ── Word (.docx, standard library) ───────────────────────────────────────────

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx(data: bytes, limits: Limits) -> Extracted:
    archive = zipfile.ZipFile(io.BytesIO(data))
    body = ElementTree.fromstring(archive.read("word/document.xml")).find(f"{_W}body")
    headings = _heading_levels(archive)
    out: list[str] = []
    for block in [] if body is None else list(body):
        if block.tag == f"{_W}p":
            line = _paragraph(block, headings)
            if line.startswith("- ") and out and out[-1].startswith("- "):
                out[-1] += "\n" + line  # keep a list together
            elif line:
                out.append(line)
        elif block.tag == f"{_W}tbl":
            out.append(_word_table(block))
    return Extracted("\n\n".join(out))


def _heading_levels(archive: zipfile.ZipFile) -> dict[str, int]:
    """Style id -> heading level. Ids are localized ("Heading1", "1", "Ueberschrift1"),
    so the style's own name ("heading 1", "Title") decides."""
    try:
        styles = ElementTree.fromstring(archive.read("word/styles.xml"))
    except KeyError:
        return {}
    levels = {}
    for style in styles.iter(f"{_W}style"):
        name_el = style.find(f"{_W}name")
        name = (name_el.get(f"{_W}val") or "").lower() if name_el is not None else ""
        if name == "title":
            levels[style.get(f"{_W}styleId", "")] = 1
        elif m := re.fullmatch(r"heading ([1-6])", name):
            levels[style.get(f"{_W}styleId", "")] = int(m.group(1))
    return levels


def _runs_text(element) -> str:
    parts = []
    for node in element.iter():
        if node.tag == f"{_W}t":
            parts.append(node.text or "")
        elif node.tag == f"{_W}tab":
            parts.append("\t")
        elif node.tag in (f"{_W}br", f"{_W}cr"):
            parts.append("\n")
    return "".join(parts)


def _paragraph(p, headings: dict[str, int]) -> str:
    text = _runs_text(p).strip()
    if not text:
        return ""
    props = p.find(f"{_W}pPr")
    style = props.find(f"{_W}pStyle") if props is not None else None
    level = headings.get(style.get(f"{_W}val", "")) if style is not None else None
    if level:
        return f"{'#' * level} {' '.join(text.split())}"
    if props is not None and props.find(f"{_W}numPr") is not None:
        return f"- {text}"
    return text


def _word_table(tbl) -> str:
    rows = []
    for tr in tbl.iter(f"{_W}tr"):
        cells = [" ".join(_runs_text(tc).split()) for tc in tr.findall(f"{_W}tc")]
        if any(cells):
            rows.append(cells)
    return _markdown_table(rows)


# ── Excel (.xlsx, openpyxl) ──────────────────────────────────────────────────


def _xlsx(data: bytes, limits: Limits) -> Extracted:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DocumentError(
            "needs_extra", 'Reading Excel needs: pip install "cairnkit[documents]"'
        ) from exc
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out, warnings, read = [], [], 0
    try:
        for sheet in book.worksheets:
            if sheet.sheet_state != "visible":
                warnings.append(f"Hidden sheet skipped: {sheet.title}")
                continue
            rows, cut = [], False
            for values in sheet.iter_rows(values_only=True):
                cells = [_cell(v) for v in values]
                if not any(cells):
                    continue
                if len(rows) >= limits.max_sheet_rows:
                    cut = True
                    break
                rows.append(cells)
            read += 1
            if not rows:
                continue
            if cut:
                warnings.append(
                    f"Sheet {sheet.title}: only the first {limits.max_sheet_rows:,} rows were read."
                )
            out.append(f"## {sheet.title}\n\n{_markdown_table(rows)}")
    finally:
        book.close()
    return Extracted("\n\n".join(out), read, warnings)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat(" ")
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return " ".join(str(value).split())


# ── shared ───────────────────────────────────────────────────────────────────


def _markdown_table(rows: list[list[str]]) -> str:
    width = max((max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows), default=0)
    if not width:
        return ""
    lines = []
    for index, row in enumerate(rows):
        cells = [c.replace("|", "\\|") for c in (row + [""] * width)[:width]]
        lines.append("| " + " | ".join(cells) + " |")
        if index == 0:
            lines.append("|" + " --- |" * width)
    return "\n".join(lines)


register_extractor("text", _text)
register_extractor("pdf", _pdf)
register_extractor("docx", _docx)
register_extractor("xlsx", _xlsx)
