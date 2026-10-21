"""Document extraction: every fixture is built here, so the tests need no files."""

import io
import zipfile
from datetime import date, datetime

import pytest

from cairnkit import documents
from cairnkit.documents import (
    DocumentError,
    Extracted,
    Limits,
    detect_kind,
    extract_text,
    register_extractor,
)

# ── builders ─────────────────────────────────────────────────────────────────


def make_pdf(pages: list[str | None]) -> bytes:
    """A small valid PDF; ``None`` makes a page with no text (like a scan)."""
    objects: list[bytes] = []
    kids = []
    font_id = 3 + 2 * len(pages)
    for index, text in enumerate(pages):
        page_id, content_id = 3 + 2 * index, 4 + 2 * index
        kids.append(f"{page_id} 0 R")
        stream = b"" if text is None else f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    body = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode(),
        *objects,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for number, obj in enumerate(body, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(body) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(body) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def make_docx(body_xml: str, styles_xml: str | None = None) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", f"<w:document {W}><w:body>{body_xml}</w:body></w:document>")
        if styles_xml is not None:
            z.writestr("word/styles.xml", f"<w:styles {W}>{styles_xml}</w:styles>")
    return buf.getvalue()


def para(text: str, style: str | None = None, listed: bool = False) -> str:
    props = ""
    if style or listed:
        props = "<w:pPr>"
        props += f'<w:pStyle w:val="{style}"/>' if style else ""
        props += '<w:numPr><w:numId w:val="1"/></w:numPr>' if listed else ""
        props += "</w:pPr>"
    return f"<w:p>{props}<w:r><w:t>{text}</w:t></w:r></w:p>"


def make_xlsx(fill) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    fill(book)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


# ── detection and limits ─────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("data", "code"),
    [
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100, "old_format"),
        (b"\x89PNG\r\n\x1a\n" + b"\x00" * 100, "image"),
        (b"RIFF\x00\x00\x00\x00WEBPVP8 ", "image"),
        (b"\x00\x01\x02binary\x00", "unsupported"),
        (b"PK\x03\x04broken", "unreadable"),
    ],
)
def test_detection_refuses_what_it_cannot_read(data, code):
    with pytest.raises(DocumentError) as exc:
        extract_text(data, "file.bin")
    assert exc.value.code == code


def test_a_zip_that_is_not_office_is_unsupported():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("readme.txt", "hello")
    assert pytest.raises(DocumentError, extract_text, buf.getvalue()).value.code == "unsupported"


def test_size_limits():
    with pytest.raises(DocumentError) as exc:
        extract_text(b"x" * 11, "a.txt", limits=Limits(max_bytes=10))
    assert exc.value.code == "too_large"
    assert pytest.raises(DocumentError, extract_text, b"", "a.txt").value.code == "empty"
    bomb = make_docx(para("A" * 5000))
    with pytest.raises(DocumentError) as exc:
        extract_text(bomb, "a.docx", limits=Limits(max_unzipped_bytes=1000))
    assert exc.value.code == "too_large"


def test_long_text_is_cut_at_a_line_and_flagged():
    data = "\n".join(f"line {i}" for i in range(2000)).encode()
    doc = extract_text(data, "notes.txt", limits=Limits(max_chars=1000))
    assert (
        doc.truncated
        and doc.chars <= 1000
        and doc.text.endswith(("0", "1", "2", "3", "4", "5", "6", "7", "8", "9"))
    )
    assert any("first 1,000 characters" in w for w in doc.warnings)


# ── text ─────────────────────────────────────────────────────────────────────


def test_plain_text_utf8_utf16_and_legacy():
    assert extract_text("Café menu\r\n\r\n\r\n\r\nItem".encode(), "m.txt").text == (
        "Café menu\n\nItem"
    )
    assert extract_text("Opening times".encode("utf-16"), "t.txt").text == "Opening times"
    legacy = extract_text("Café".encode("cp1252"), "old.txt")
    assert legacy.text == "Café" and legacy.warnings
    assert detect_kind(b"name,handicap\nAvery,15\n", "players.csv") == "text"


# ── PDF ──────────────────────────────────────────────────────────────────────


def test_pdf_pages_and_scanned_pages():
    doc = extract_text(make_pdf(["House rules", None, "Guest policy"]), "rules.pdf")
    assert doc.kind == "pdf" and doc.parts == 3
    assert "[Page 1]\nHouse rules" in doc.text and "[Page 3]\nGuest policy" in doc.text
    assert any("1 of 3 pages had no text" in w for w in doc.warnings)


def test_pdf_without_any_text_is_reported_as_a_scan():
    with pytest.raises(DocumentError) as exc:
        extract_text(make_pdf([None, None]), "scan.pdf")
    assert exc.value.code == "no_text"


def test_pdf_page_limit():
    doc = extract_text(make_pdf(["One", "Two", "Three"]), "a.pdf", limits=Limits(max_pages=2))
    assert "Three" not in doc.text and any("first 2 of 3 pages" in w for w in doc.warnings)


def test_password_protected_pdf():
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(make_pdf(["Secret"]))))
    writer.encrypt("open-sesame", algorithm="RC4-128")
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(DocumentError) as exc:
        extract_text(buf.getvalue(), "locked.pdf")
    assert exc.value.code == "encrypted"


def test_damaged_pdf_is_unreadable():
    with pytest.raises(DocumentError) as exc:
        extract_text(b"%PDF-1.4\nnot really a pdf", "bad.pdf")
    assert exc.value.code in ("unreadable", "no_text")


# ── Word ─────────────────────────────────────────────────────────────────────


def test_docx_headings_lists_tables_and_breaks():
    styles = (
        '<w:style w:styleId="1"><w:name w:val="heading 1"/></w:style>'  # a localized id
        '<w:style w:styleId="Heading2"><w:name w:val="heading 2"/></w:style>'
    )
    body = (
        para("Guest Policy", style="1")
        + para("Who may play", style="Heading2")
        + para("Members may bring two guests a month.")
        + para("Guests pay the guest fee", listed=True)
        + para("Guests check in at the pro shop", listed=True)
        + "<w:p><w:r><w:t>Line one</w:t><w:br/><w:t>line two</w:t></w:r></w:p>"
        + "<w:tbl><w:tr><w:tc>"
        + para("Day")
        + "</w:tc><w:tc>"
        + para("Fee")
        + "</w:tc></w:tr>"
        + "<w:tr><w:tc>"
        + para("Sat")
        + "</w:tc><w:tc>"
        + para("$95 | cart")
        + "</w:tc></w:tr>"
        + "</w:tbl>"
        + "<w:p><w:del><w:r><w:delText>deleted</w:delText></w:r></w:del></w:p>"
    )
    doc = extract_text(make_docx(body, styles), "policy.docx")
    assert doc.kind == "docx"
    assert doc.text.split("\n\n")[:4] == [
        "# Guest Policy",
        "## Who may play",
        "Members may bring two guests a month.",
        "- Guests pay the guest fee\n- Guests check in at the pro shop",
    ]
    assert "Line one\nline two" in doc.text
    assert "| Day | Fee |\n| --- | --- |\n| Sat | $95 \\| cart |" in doc.text
    assert "deleted" not in doc.text


def test_empty_docx_has_no_text():
    with pytest.raises(DocumentError) as exc:
        extract_text(make_docx(""), "blank.docx")
    assert exc.value.code == "no_text"


# ── Excel ────────────────────────────────────────────────────────────────────


def test_xlsx_sheets_values_and_hidden_sheets():
    def fill(book):
        prices = book.active
        prices.title = "Prices"
        prices.append(["Item", "Price", "Since", "Members only"])
        prices.append(["Day pass", 8.0, date(2026, 3, 1), False])
        prices.append([None, None, None, None])
        prices.append(["Lesson", 95.5, datetime(2026, 4, 2, 9, 30), True])
        events = book.create_sheet("Events")
        events.append(["Open day", datetime(2026, 11, 7)])
        hidden = book.create_sheet("Notes")
        hidden.append(["internal"])
        hidden.sheet_state = "hidden"

    doc = extract_text(make_xlsx(fill), "prices.xlsx")
    assert doc.kind == "xlsx" and doc.parts == 2
    assert (
        "## Prices\n\n| Item | Price | Since | Members only |\n| --- | --- | --- | --- |"
        in doc.text
    )
    assert "| Day pass | 8 | 2026-03-01 | FALSE |" in doc.text
    assert "| Lesson | 95.5 | 2026-04-02 09:30:00 | TRUE |" in doc.text
    assert "## Events\n\n| Open day | 2026-11-07 |" in doc.text
    assert "internal" not in doc.text and any("Hidden sheet" in w for w in doc.warnings)


def test_xlsx_row_limit():
    def fill(book):
        for i in range(30):
            book.active.append([f"row {i}"])

    doc = extract_text(make_xlsx(fill), "big.xlsx", limits=Limits(max_sheet_rows=10))
    assert "row 9" in doc.text and "row 10" not in doc.text
    assert any("first 10 rows" in w for w in doc.warnings)


# ── replaceable backends and missing extras ─────────────────────────────────


def test_a_project_can_swap_in_its_own_reader():
    previous = register_extractor("pdf", lambda data, limits: Extracted("from OCR", 1))
    try:
        assert extract_text(make_pdf([None]), "scan.pdf").text == "from OCR"
    finally:
        register_extractor("pdf", previous)


def test_missing_optional_dependency_is_explained(monkeypatch):
    import builtins

    real_import = builtins.__import__

    def no_pypdf(name, *args, **kwargs):
        if name == "pypdf":
            raise ImportError(name)
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", no_pypdf)
    with pytest.raises(DocumentError) as exc:
        documents.extract_text(make_pdf(["x"]), "a.pdf")
    assert exc.value.code == "needs_extra"
